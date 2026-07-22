#!/usr/bin/env python3
"""Replace duplicate trials 100 and 108 with screened, publication-aligned RCTs.

The source workbooks are ignored local screening artifacts.  This script keeps
the public outputs compact and participant-level, archives the superseded raw
sources, writes CSV/RDS pairs, and records publication-value audit rows.
"""

from __future__ import annotations

import csv
import math
import shutil
import subprocess
from collections import Counter
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
CLEAN = ROOT / "cleaned_data"
LOCAL = ROOT / "local" / "rct_expansion"
RAW = LOCAL / "raw_data"
PROVENANCE = LOCAL / "provenance"
SCREEN = PROVENANCE / "landing_page_recheck_2026_06_11" / "downloads"
ARCHIVE = LOCAL / "backup" / "duplicate_replacements_2026_07_22"

PT_SOURCE = SCREEN / "RCTC-02610" / "PT vs. Home ex. data.xlsx"
EXTRACT_SOURCE = SCREEN / "RCTC-06339" / "EXTRACTNOAC_FAS.xlsx"


def clean_number(value: Any) -> int | float | None:
    if value is None or value == "":
        return None
    if isinstance(value, bool):
        return int(value)
    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            return None
        return value
    try:
        return float(str(value).strip())
    except ValueError:
        return None


def difference(end: Any, baseline: Any, divisor: float = 1.0) -> float | None:
    end_num = clean_number(end)
    baseline_num = clean_number(baseline)
    if end_num is None or baseline_num is None:
        return None
    return (float(end_num) - float(baseline_num)) / divisor


def divided(value: Any, divisor: float) -> float | None:
    number = clean_number(value)
    return None if number is None else float(number) / divisor


def archive_and_install_source(trial_id: int, source: Path, superseded_names: Iterable[str]) -> None:
    trial_raw = RAW / f"trial{trial_id}"
    trial_archive = ARCHIVE / f"trial{trial_id}"
    trial_raw.mkdir(parents=True, exist_ok=True)
    trial_archive.mkdir(parents=True, exist_ok=True)
    for name in superseded_names:
        old = trial_raw / name
        if old.exists():
            destination = trial_archive / name
            if destination.exists():
                old.unlink()
            else:
                shutil.move(str(old), str(destination))
    shutil.copy2(source, trial_raw / source.name)


def read_pt_rows() -> list[tuple[Any, ...]]:
    ws = load_workbook(PT_SOURCE, read_only=True, data_only=True)["Sheet1"]
    rows = list(ws.iter_rows(min_row=3, values_only=True))
    if len(rows) != 86 or Counter(row[0] for row in rows) != Counter({1: 43, 2: 43}):
        raise ValueError("Unexpected supervised-PT source cohort or allocation counts")
    return rows


def clean_trial100() -> list[dict[str, Any]]:
    rows = read_pt_rows()
    treatment = {1: "Supervised physical therapy", 2: "Home exercise"}
    sex = {1: "Male", 2: "Female"}
    cleaned: list[dict[str, Any]] = []
    for row in rows:
        cleaned.append(
            {
                "Treatment": treatment[row[0]],
                "YP_delta_zcq_symptom_severity_6w": difference(row[53], row[28], 7.0),
                "YS_delta_zcq_physical_function_6w": difference(row[54], row[29], 5.0),
                "YS_zcq_satisfaction_6w": clean_number(row[55]),
                "YS_delta_nrs_back_pain_6w": difference(row[37], row[12]),
                "YS_delta_nrs_leg_pain_6w": difference(row[38], row[13]),
                "YS_delta_nrs_leg_numbness_6w": difference(row[39], row[14]),
                "YS_delta_walk_distance_m_6w": difference(row[61], row[35]),
                "YS_delta_daily_steps_6w": difference(row[88], row[87]),
                "X_zcq_symptom_severity_0w": divided(row[28], 7.0),
                "X_zcq_physical_function_0w": divided(row[29], 5.0),
                "X_nrs_back_pain_0w": clean_number(row[12]),
                "X_nrs_leg_pain_0w": clean_number(row[13]),
                "X_nrs_leg_numbness_0w": clean_number(row[14]),
                "X_walk_distance_m_0w": clean_number(row[35]),
                "X_daily_steps_0w": clean_number(row[87]),
                "X_age_years": clean_number(row[1]),
                "X_sex": sex.get(row[2]),
                "X_bmi": clean_number(row[3]),
                "X_symptom_duration_months": clean_number(row[4]),
            }
        )
    level_order = {"Home exercise": 0, "Supervised physical therapy": 1}
    cleaned.sort(key=lambda row: level_order[row["Treatment"]])
    return cleaned


def derive_noac(row: tuple[Any, ...]) -> str | None:
    labels = ("Dabigatran", "Rivaroxaban", "Apixaban", "Edoxaban")
    for label, value in zip(labels, row[4:8]):
        if value not in (None, "", 0, "0"):
            return label
    return None


def clean_trial108() -> list[dict[str, Any]]:
    ws = load_workbook(EXTRACT_SOURCE, read_only=True, data_only=True)["Data"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    if len(rows) != 218 or Counter(row[0] for row in rows) != Counter({"Placebo": 112, "TXA": 106}):
        raise ValueError("Unexpected EXTRACT-NOAC full-analysis cohort or allocation counts")
    cleaned: list[dict[str, Any]] = []
    for row in rows:
        cleaned.append(
            {
                "Treatment": row[0],
                "YP_any_oral_bleeding_7d": clean_number(row[31]),
                "YS_oral_bleeding_count_7d": clean_number(row[32]),
                "YS_clinically_relevant_oral_bleeding_7d": clean_number(row[35]),
                "YS_clinically_relevant_oral_bleeding_count_7d": clean_number(row[36]),
                "YS_early_oral_bleeding_7d": clean_number(row[39]),
                "YS_early_oral_bleeding_count_7d": clean_number(row[40]),
                "YS_delayed_oral_bleeding_7d": clean_number(row[41]),
                "YS_delayed_oral_bleeding_count_7d": clean_number(row[42]),
                "YS_unplanned_medical_contact_7d": clean_number(row[61]),
                "YS_thrombotic_event_7d": clean_number(row[67]),
                "X_age_years": clean_number(row[2]),
                "X_sex": row[1],
                "X_noac_type": derive_noac(row),
                "X_alcohol_units_per_day": clean_number(row[8]),
                "X_smoking_status": row[9],
                "X_chronic_heart_failure": clean_number(row[10]),
                "X_hypertension": clean_number(row[11]),
                "X_diabetes": clean_number(row[12]),
                "X_history_stroke": clean_number(row[13]),
                "X_chads_vasc_score": clean_number(row[16]),
                "X_noac_indication": row[17],
                "X_time_last_noac_to_extraction_days": clean_number(row[18]),
                "X_number_extracted_teeth": clean_number(row[20]),
            }
        )
    cleaned.sort(key=lambda row: 0 if row["Treatment"] == "Placebo" else 1)
    return cleaned


def write_csv(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def write_rds(trial_id: int, treatment_levels: list[str]) -> None:
    csv_path = CLEAN / f"trial{trial_id}.csv"
    rds_path = CLEAN / f"trial{trial_id}.rds"
    levels = ",".join(f'"{value}"' for value in treatment_levels)
    expression = (
        f'd <- read.csv("{csv_path}", check.names=FALSE, na.strings=c("", "NA")); '
        f'd$Treatment <- factor(d$Treatment, levels=c({levels})); '
        f'saveRDS(d, "{rds_path}")'
    )
    subprocess.run(["Rscript", "-e", expression], check=True)


def audit_rows(trial100: list[dict[str, Any]], trial108: list[dict[str, Any]]) -> list[dict[str, Any]]:
    pt_by_arm = {arm: [row for row in trial100 if row["Treatment"] == arm] for arm in ("Home exercise", "Supervised physical therapy")}
    pt_responder = {}
    for arm, rows in pt_by_arm.items():
        pt_responder[arm] = sum(
            row["YP_delta_zcq_symptom_severity_6w"] is not None
            and row["YP_delta_zcq_symptom_severity_6w"] <= -0.36
            for row in rows
        )
    pt_rates = {arm: 100 * value / 43 for arm, value in pt_responder.items()}
    pt_difference = pt_rates["Supervised physical therapy"] - pt_rates["Home exercise"]
    if pt_responder != {"Home exercise": 14, "Supervised physical therapy": 27} or abs(pt_difference - 30.2) > 0.05:
        raise ValueError("Supervised-PT primary responder result does not reproduce the publication")

    ex_by_arm = {arm: [row for row in trial108 if row["Treatment"] == arm] for arm in ("Placebo", "TXA")}
    ex_events = {arm: int(sum(row["YP_any_oral_bleeding_7d"] or 0 for row in rows)) for arm, rows in ex_by_arm.items()}
    ex_counts = {arm: len(rows) for arm, rows in ex_by_arm.items()}
    if ex_counts != {"Placebo": 112, "TXA": 106} or ex_events != {"Placebo": 32, "TXA": 28}:
        raise ValueError("EXTRACT-NOAC primary result does not reproduce the publication")

    columns = [
        "Trial_ID", "audit_source", "audit_type", "candidate_id", "Original_Trial_ID",
        "outcome_variable", "outcome_role", "arm", "statistic", "paper_value",
        "paper_precision", "cleaned_value", "absolute_diff", "tolerance", "status",
        "paper_source", "notes",
    ]
    result: list[dict[str, Any]] = []

    def add(trial_id: int, candidate_id: str, outcome: str, role: str, arm: str, statistic: str,
            paper: float, precision: float, cleaned: float, source: str, notes: str) -> None:
        tolerance = max(abs(paper) * 0.05, precision)
        diff = abs(cleaned - paper)
        result.append(dict(zip(columns, [
            str(trial_id), "replacement_trials100_108_audit.csv", "strict_publication_reproducibility",
            candidate_id, None, outcome, role, arm, statistic, paper, precision, cleaned,
            diff, tolerance, "pass" if diff <= tolerance else "fail", source, notes,
        ])))

    add(100, "RCTC-02610", "Treatment", "allocation", "Home exercise", "n", 43, 1, 43,
        "Minetama et al. Spine Journal 2019 abstract/results", "Randomized allocation count.")
    add(100, "RCTC-02610", "Treatment", "allocation", "Supervised physical therapy", "n", 43, 1, 43,
        "Minetama et al. Spine Journal 2019 abstract/results", "Randomized allocation count.")
    add(100, "RCTC-02610", "YP_delta_zcq_symptom_severity_6w", "primary", "Home exercise",
        "mcid_responder_percent_itt", 32.6, 0.1, pt_rates["Home exercise"],
        "Minetama et al. Spine Journal 2019 results", "MCID responder rate; missing 6-week outcome retained as nonresponder in the ITT denominator.")
    add(100, "RCTC-02610", "YP_delta_zcq_symptom_severity_6w", "primary", "Supervised physical therapy",
        "mcid_responder_percent_itt", 62.8, 0.1, pt_rates["Supervised physical therapy"],
        "Minetama et al. Spine Journal 2019 results", "MCID responder rate; missing 6-week outcome retained as nonresponder in the ITT denominator.")
    add(100, "RCTC-02610", "YP_delta_zcq_symptom_severity_6w", "primary", "Supervised physical therapy minus home exercise",
        "mcid_responder_percentage_point_difference", 30.2, 0.1, pt_difference,
        "Minetama et al. Spine Journal 2019 abstract/results", "Published primary between-arm responder-rate difference.")

    for arm, paper_n, paper_events, paper_percent in (("Placebo", 112, 32, 28.6), ("TXA", 106, 28, 26.4)):
        add(108, "RCTC-06339", "Treatment", "allocation", arm, "n", paper_n, 1, ex_counts[arm],
            "Ockerman et al. PLOS Medicine 2021 abstract/results", "Full analysis set allocation count.")
        add(108, "RCTC-06339", "YP_any_oral_bleeding_7d", "primary", arm, "event_count", paper_events, 1, ex_events[arm],
            "Ockerman et al. PLOS Medicine 2021 Table 2", "Patients with any oral post-extraction bleeding through day 7.")
        add(108, "RCTC-06339", "YP_any_oral_bleeding_7d", "primary", arm, "percent", paper_percent, 0.1,
            100 * ex_events[arm] / ex_counts[arm], "Ockerman et al. PLOS Medicine 2021 Table 2",
            "Patients with any oral post-extraction bleeding through day 7.")
    return result


def write_audit(rows: list[dict[str, Any]]) -> None:
    path = PROVENANCE / "replacement_trials100_108_audit.csv"
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    if not PT_SOURCE.exists() or not EXTRACT_SOURCE.exists():
        raise FileNotFoundError("Screened replacement source workbook is missing")
    trial100 = clean_trial100()
    trial108 = clean_trial108()
    archive_and_install_source(
        100,
        PT_SOURCE,
        ["coagulation_data_clean_R Open data.xlsx"],
    )
    archive_and_install_source(
        108,
        EXTRACT_SOURCE,
        ["ELAIA-1_deidentified_data_10-6-2020 (1).csv", "ELAIA-1_data_dictionary_10-6-2020 (1).csv"],
    )
    write_csv(CLEAN / "trial100.csv", trial100)
    write_csv(CLEAN / "trial108.csv", trial108)
    write_rds(100, ["Home exercise", "Supervised physical therapy"])
    write_rds(108, ["Placebo", "TXA"])
    audits = audit_rows(trial100, trial108)
    write_audit(audits)
    print(f"Wrote trial100: {len(trial100)} rows, {len(trial100[0])} columns")
    print(f"Wrote trial108: {len(trial108)} rows, {len(trial108[0])} columns")
    print(f"Wrote {len(audits)} strict publication-audit rows")


if __name__ == "__main__":
    main()
