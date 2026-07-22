#!/usr/bin/env python3
"""Screen all public trials for metadata collisions and participant-data duplicates."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import unicodedata
from collections import Counter, defaultdict
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Callable

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
DEFAULT_CSV = ROOT / "local/rct_expansion/provenance/duplicate_database_audit.csv"
DEFAULT_JSON = ROOT / "local/rct_expansion/provenance/duplicate_database_audit_summary.json"
MISSING = {"", "na", "nan", "n/a", "null", "none"}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-csv", type=Path, default=DEFAULT_CSV)
    parser.add_argument("--output-json", type=Path, default=DEFAULT_JSON)
    return parser.parse_args()


def normalize_text(value: object) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = text.encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]+", "", text)


def normalize_value(value: object) -> str:
    text = str(value or "").strip()
    if text.lower() in MISSING:
        return "<NA>"
    try:
        number = round(Decimal(text), 12)
        if number == 0:
            number = Decimal(0)
        return format(number.normalize(), "f")
    except InvalidOperation:
        return normalize_text(text)


def extract_doi(value: object) -> str:
    match = re.search(r"10\.\d{4,9}/[^\s?#]+", str(value or "").lower())
    return match.group(0).rstrip(".,;)") if match else ""


def extract_registry_or_name(value: object) -> str:
    text = str(value or "").upper().strip()
    patterns = [
        r"NCT\d{8}", r"ISRCTN\d+", r"UMIN(?:-?CTR)?\s*[-:]?\s*\d+",
        r"ACTRN\d+", r"DRKS\d+", r"CHICTR[-A-Z0-9]+",
        r"CTRI[/A-Z0-9-]+", r"TCTR\d+", r"PACTR\d+", r"EUCTR[-A-Z0-9]+",
    ]
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            return re.sub(r"[^A-Z0-9]", "", match.group(0))
    normalized = normalize_text(text)
    if normalized in {"", "na", "none", "notreported", "notapplicable"}:
        return ""
    return normalized if len(normalized) >= 6 else ""


def file_hash(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def trial_number(path: Path) -> int:
    match = re.fullmatch(r"trial(\d+)\.(?:csv|rds)", path.name)
    if not match:
        raise ValueError(f"Unexpected trial filename: {path}")
    return int(match.group(1))


def read_metadata() -> dict[int, dict[str, object]]:
    workbook = load_workbook(ROOT / "meta_data.xlsx", data_only=True, read_only=True)
    sheet = workbook["Sheet1"]
    rows = list(sheet.iter_rows(values_only=True))
    headers = [str(value or "") for value in rows[0][:18]]
    return {int(row[0]): dict(zip(headers, row[:18])) for row in rows[1:] if row[0] is not None}


def read_tables() -> dict[int, dict[str, object]]:
    tables: dict[int, dict[str, object]] = {}
    for path in sorted((ROOT / "cleaned_data").glob("trial*.csv"), key=trial_number):
        with path.open(newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle)
            rows = list(reader)
            columns = list(reader.fieldnames or [])
        tables[trial_number(path)] = {"path": path, "columns": columns, "rows": rows}
    return tables


def duplicate_groups(
    metadata: dict[int, dict[str, object]], column: str, normalizer: Callable[[object], str]
) -> dict[str, list[int]]:
    groups: dict[str, list[int]] = defaultdict(list)
    for identifier, row in metadata.items():
        key = normalizer(row.get(column))
        if key:
            groups[key].append(identifier)
    return {key: values for key, values in groups.items() if len(values) > 1}


def pairs_from_groups(groups: dict[str, list[int]], label: str) -> dict[tuple[int, int], list[str]]:
    signals: dict[tuple[int, int], list[str]] = defaultdict(list)
    for ids in groups.values():
        for position, left in enumerate(ids):
            for right in ids[position + 1 :]:
                signals[(min(left, right), max(left, right))].append(label)
    return signals


def merge_pair_signals(*maps: dict[tuple[int, int], list[str]]) -> dict[tuple[int, int], list[str]]:
    merged: dict[tuple[int, int], list[str]] = defaultdict(list)
    for mapping in maps:
        for pair, values in mapping.items():
            merged[pair].extend(values)
    return merged


def hash_pairs(extension: str) -> set[tuple[int, int]]:
    groups: dict[str, list[int]] = defaultdict(list)
    for path in (ROOT / "cleaned_data").glob(f"trial*.{extension}"):
        groups[file_hash(path)].append(trial_number(path))
    result: set[tuple[int, int]] = set()
    for ids in groups.values():
        if len(ids) < 2:
            continue
        for position, left in enumerate(ids):
            result.update((min(left, right), max(left, right)) for right in ids[position + 1 :])
    return result


def vector_profiles(tables: dict[int, dict[str, object]]) -> dict[int, dict[str, object]]:
    profiles: dict[int, dict[str, object]] = {}
    for identifier, table in tables.items():
        ordered: dict[str, list[str]] = defaultdict(list)
        unordered: dict[str, list[str]] = defaultdict(list)
        rows = table["rows"]
        for column in table["columns"]:
            values = [normalize_value(row[column]) for row in rows]
            if len(set(values)) <= 1:
                continue
            ordered_hash = hashlib.sha256("\x1f".join(values).encode()).hexdigest()
            unordered_hash = hashlib.sha256("\x1f".join(sorted(values)).encode()).hexdigest()
            ordered[ordered_hash].append(column)
            unordered[unordered_hash].append(column)
        profiles[identifier] = {
            "n": len(rows),
            "p": len(table["columns"]),
            "ordered": ordered,
            "unordered": unordered,
        }
    return profiles


def shared_named_match(tables: dict[int, dict[str, object]], left: int, right: int) -> tuple[int, float | None]:
    a, b = tables[left], tables[right]
    if len(a["rows"]) != len(b["rows"]):
        return 0, None
    shared = [name for name in a["columns"] if name in set(b["columns"])]
    nonconstant = [
        name for name in shared
        if len({normalize_value(row[name]) for row in a["rows"]}) > 1
        and len({normalize_value(row[name]) for row in b["rows"]}) > 1
    ]
    if not nonconstant:
        return 0, None
    matches = sum(
        all(normalize_value(a_row[name]) == normalize_value(b_row[name]) for name in nonconstant)
        for a_row, b_row in zip(a["rows"], b["rows"])
    )
    return len(nonconstant), matches / len(a["rows"])


def numeric_profiles(tables: dict[int, dict[str, object]]) -> dict[int, list[tuple[str, Counter[str], int]]]:
    profiles: dict[int, list[tuple[str, Counter[str], int]]] = {}
    for identifier, table in tables.items():
        columns: list[tuple[str, Counter[str], int]] = []
        for column in table["columns"]:
            values: list[Decimal] = []
            for row in table["rows"]:
                text = str(row[column] or "").strip().lower()
                if text in MISSING:
                    continue
                try:
                    values.append(Decimal(text))
                except InvalidOperation:
                    continue
            if len(values) < 20:
                continue
            normalized = [format(value.normalize(), "f") for value in values]
            distinct = len(set(normalized))
            fractional_share = sum(value != value.to_integral_value() for value in values) / len(values)
            if distinct >= 20 and distinct / len(values) >= 0.20 and fractional_share >= 0.25:
                columns.append((column, Counter(normalized), len(normalized)))
        profiles[identifier] = columns
    return profiles


def high_cardinality_overlap(
    numeric: dict[int, list[tuple[str, Counter[str], int]]], left: int, right: int
) -> list[tuple[float, str, str]]:
    possible: list[tuple[float, str, str]] = []
    for left_name, left_values, left_count in numeric[left]:
        for right_name, right_values, right_count in numeric[right]:
            overlap = sum((left_values & right_values).values()) / min(left_count, right_count)
            if overlap >= 0.80:
                possible.append((overlap, left_name, right_name))
    selected: list[tuple[float, str, str]] = []
    used_left: set[str] = set()
    used_right: set[str] = set()
    for match in sorted(possible, reverse=True):
        if match[1] in used_left or match[2] in used_right:
            continue
        selected.append(match)
        used_left.add(match[1])
        used_right.add(match[2])
    return selected


def main() -> None:
    args = parse_args()
    metadata = read_metadata()
    tables = read_tables()
    if sorted(metadata) != list(range(1, 126)) or sorted(tables) != list(range(1, 126)):
        raise RuntimeError("Expected complete trial1-trial125 metadata and CSV coverage")

    metadata_pairs = merge_pair_signals(
        pairs_from_groups(duplicate_groups(metadata, "Trial Number/Name", extract_registry_or_name), "registry_or_trial_name"),
        pairs_from_groups(duplicate_groups(metadata, "Paper Link", extract_doi), "paper_doi"),
        pairs_from_groups(duplicate_groups(metadata, "Paper Name", normalize_text), "paper_title"),
    )
    exact_csv = hash_pairs("csv")
    exact_rds = hash_pairs("rds")
    vectors = vector_profiles(tables)
    numeric = numeric_profiles(tables)

    candidates: list[dict[str, object]] = []
    all_pairs = [(left, right) for left in range(1, 126) for right in range(left + 1, 126)]
    for left, right in all_pairs:
        left_profile, right_profile = vectors[left], vectors[right]
        same_n = left_profile["n"] == right_profile["n"]
        ordered_shared = len(set(left_profile["ordered"]) & set(right_profile["ordered"])) if same_n else 0
        unordered_shared = len(set(left_profile["unordered"]) & set(right_profile["unordered"])) if same_n else 0
        named_columns, named_match_rate = shared_named_match(tables, left, right)
        sample_ratio = min(left_profile["n"], right_profile["n"]) / max(left_profile["n"], right_profile["n"])
        numeric_overlap = high_cardinality_overlap(numeric, left, right) if sample_ratio >= 0.50 else []
        metadata_signals = sorted(set(metadata_pairs.get((left, right), [])))
        exact_csv_match = (left, right) in exact_csv
        exact_rds_match = (left, right) in exact_rds

        strong_data = (
            exact_csv_match or exact_rds_match or ordered_shared >= 4
            or unordered_shared >= 5 or (named_columns >= 3 and named_match_rate == 1)
        )
        moderate_data = ordered_shared >= 2 or unordered_shared >= 3 or len(numeric_overlap) >= 2
        if strong_data and metadata_signals:
            classification = "confirmed_duplicate"
        elif strong_data:
            classification = "data_duplicate_review"
        elif metadata_signals:
            classification = "metadata_collision_review"
        elif moderate_data:
            classification = "data_overlap_review"
        else:
            continue

        evidence_parts = []
        if metadata_signals:
            evidence_parts.append("metadata=" + ";".join(metadata_signals))
        if exact_csv_match:
            evidence_parts.append("exact CSV hash")
        if exact_rds_match:
            evidence_parts.append("exact RDS hash")
        if ordered_shared:
            evidence_parts.append(f"{ordered_shared} ordered column-vector fingerprints")
        if unordered_shared:
            evidence_parts.append(f"{unordered_shared} unordered column-vector fingerprints")
        if named_match_rate is not None:
            evidence_parts.append(f"{named_columns} shared named vectors; row match={named_match_rate:.3f}")
        if numeric_overlap:
            evidence_parts.append(
                "numeric multiset overlap="
                + ";".join(f"{a}:{b}:{rate:.3f}" for rate, a, b in numeric_overlap)
            )
        candidates.append({
            "trial_left": left,
            "trial_right": right,
            "classification": classification,
            "metadata_signals": ";".join(metadata_signals),
            "left_rows": left_profile["n"],
            "right_rows": right_profile["n"],
            "exact_csv_hash": exact_csv_match,
            "exact_rds_hash": exact_rds_match,
            "ordered_shared_vectors": ordered_shared,
            "unordered_shared_vectors": unordered_shared,
            "shared_named_nonconstant_columns": named_columns,
            "shared_named_row_match_rate": named_match_rate,
            "high_cardinality_numeric_overlap_count": len(numeric_overlap),
            "evidence": " | ".join(evidence_parts),
        })

    candidates.sort(key=lambda row: (str(row["classification"]), int(row["trial_left"]), int(row["trial_right"])))
    fields = [
        "trial_left", "trial_right", "classification", "metadata_signals", "left_rows", "right_rows",
        "exact_csv_hash", "exact_rds_hash", "ordered_shared_vectors", "unordered_shared_vectors",
        "shared_named_nonconstant_columns", "shared_named_row_match_rate",
        "high_cardinality_numeric_overlap_count", "evidence",
    ]
    args.output_csv.parent.mkdir(parents=True, exist_ok=True)
    with args.output_csv.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(candidates)

    counts = Counter(str(row["classification"]) for row in candidates)
    summary = {
        "trial_count": len(tables),
        "metadata_row_count": len(metadata),
        "candidate_count": len(candidates),
        "classification_counts": dict(sorted(counts.items())),
        "high_confidence_duplicate_count": counts["confirmed_duplicate"] + counts["data_duplicate_review"],
        "metadata_collision_count": counts["metadata_collision_review"],
        "outputs": {"csv": str(args.output_csv), "json": str(args.output_json)},
    }
    args.output_json.write_text(json.dumps(summary, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(summary, indent=2))
    for row in candidates:
        print(json.dumps(row, default=str))


if __name__ == "__main__":
    main()
